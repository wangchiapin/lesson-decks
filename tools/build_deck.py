import re, sys, os, importlib.util
import opencc
# 用法: python3 tools/build_deck.py <topics檔.py> <輸出資料夾>  (輸出檔名: 主題名_教材.pptx)
spec=importlib.util.spec_from_file_location('topics_mod',sys.argv[1]); mod=importlib.util.module_from_spec(spec); spec.loader.exec_module(mod)
T=mod.T; OUT=sys.argv[2]
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
t2s=opencc.OpenCC('t2s')
FONT="澳聲通拼音庫楷-繁"
BLUE,GREEN,BORDER=RGBColor(0xDA,0xE3,0xF3),RGBColor(0xE2,0xF0,0xD9),RGBColor(0x37,0x56,0x23)

def box(slide,x,y,w,h,fill,lines=None,lang="zh-TW",size=40):
    sh=slide.shapes.add_textbox(Emu(x),Emu(y),Emu(w),Emu(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=fill
    sh.line.color.rgb=BORDER; sh.line.width=Emu(19050)
    tf=sh.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    bp=tf._txBody.find(qn('a:bodyPr'))
    for c in list(bp): bp.remove(c)
    etree.SubElement(bp,qn('a:noAutofit'))
    for i,t in enumerate(lines or [""]):
        para=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        para.alignment=PP_ALIGN.CENTER; para.line_spacing=1.5
        if t:
            r=para.add_run(); r.text=t; r.font.size=Pt(size); r.font.color.rgb=RGBColor(0,0,0)
            rPr=r._r.get_or_add_rPr(); rPr.set('lang',lang)
            for tag in ('a:latin','a:ea','a:cs'): etree.SubElement(rPr,qn(tag)).set('typeface',FONT)

def make(prs,lines,notes,lang,size=40):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    box(s,4036741,147106,8053659,6588974,BLUE,lines,lang,size)
    box(s,217798,2823882,3718582,3912197,GREEN)
    ph=s.shapes.add_shape(1,Emu(217798),Emu(147106),Emu(3718582),Emu(2600000))
    ph.fill.solid(); ph.fill.fore_color.rgb=RGBColor(0xF2,0xF2,0xF2)
    ph.line.color.rgb=RGBColor(0xA6,0xA6,0xA6); ph.line.dash_style=4
    ph.text_frame.text="（在這裡放圖片）"
    s.notes_slide.notes_text_frame.text=notes

os.makedirs(OUT,exist_ok=True)
problems=[]
for i,t in enumerate(T,1):
    nm=t['name']
    # checks
    for l in t['lines']:
        if len(l)>14: problems.append(f"{nm}: line too long ({len(l)}): {l}")
    if len(t['lines'])>8 or not (10<=len(t['p2'])<=13): problems.append(f"{nm}: p2 count {len(t['p2'])}")
    for v in t['vocab']:
        w=v.split('|')[0]
        if len(v.split('|'))!=3: problems.append(f"{nm}: bad vocab {v}")
        key=w.replace('…','')
        if '…' not in w and w not in ''.join(t['lines']): problems.append(f"{nm}: vocab not in text: {w}")
    for q in t['p2']:
        if len(q)!=3: problems.append(f"{nm}: bad p2 {q}")
    # vocab lines
    vs=[f"{v.split('|')[0]}：{v.split('|')[1]} / {v.split('|')[2]}" for v in t['vocab']]
    vl=["  //  ".join(vs[k:k+3]) for k in range(0,len(vs),3)]
    if t['extra']: vl.append(t['extra'])
    p1="\n".join(vl)+"\n\n問答："+t['qa']+"\n補充："+t['supp']+"\n\n西班牙文："+t['es']+"\n英文："+t['en']
    zh_es="\n".join(f"{n}.{q[0]} {q[1]}" for n,q in enumerate(t['p2'],1))
    en="\n".join(f"{n}.{q[2]}" for n,q in enumerate(t['p2'],1))
    p2=zh_es+"\n\n"+en
    simp=[t2s.convert(l).replace('萤幕','荧幕').replace('著','着').replace('帐','账').replace('写著','写着') for l in t['lines']]
    prs=Presentation(); prs.slide_width=Emu(12192000); prs.slide_height=Emu(6858000)
    SZ=36 if len(t['lines'])>7 else 40
    make(prs,t['lines'],p1,"zh-TW",SZ); make(prs,None,p2,"zh-TW")
    make(prs,simp,p1,"zh-CN",SZ); make(prs,None,p2,"zh-CN")
    # variation selector check
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame and re.search('[\U000E0100-\U000E01EF]',sh.text_frame.text): problems.append(nm+": VS in slide")
        if re.search('[\U000E0100-\U000E01EF]',s.notes_slide.notes_text_frame.text): problems.append(nm+": VS in notes")
    prs.save(f"{OUT}/{nm}_教材.pptx")
    print(nm, '| simp:', ' / '.join(simp))
print("PROBLEMS:"); print("\n".join(problems) or "none")
