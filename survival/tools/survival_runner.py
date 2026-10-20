
import sys, importlib.util
from pptx import Presentation
import os
W0 = os.path.dirname(os.path.abspath(__file__)) + '/'
spec = importlib.util.spec_from_file_location('topic', sys.argv[1])
T = importlib.util.module_from_spec(spec); spec.loader.exec_module(T)
p = Presentation(W0 + 'survival_template.pptx')
sl = list(p.slides)
import copy, random, re
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE
from opencc import OpenCC

cc = OpenCC('t2s')
S = cc.convert
W = W0

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'

def shape(slide, name):
    return [s for s in slide.shapes if s.name == name][0]

# ---- template paragraph for vocab (zh run + es run) from slide 2
vtpl = copy.deepcopy(shape(sl[1], '文字方塊 8')._element.txBody.findall(A + 'p')[0])

def set_vocab(slide, items):
    tb = shape(slide, '文字方塊 8')._element.txBody
    for para in tb.findall(A + 'p'):
        tb.remove(para)
    for zh, es in items:
        para = copy.deepcopy(vtpl)
        runs = para.findall(A + 'r')
        runs[0].find(A + 't').text = S(zh)
        runs[1].find(A + 't').text = (' ' + es) if zh[-1].isascii() else es
        tb.append(para)

def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

def photo_placeholder(slide, hint):
    for old in [x for x in slide.shapes if x.name == '圖片位置']:
        old._element.getparent().remove(old._element)
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(203200), Emu(152400), Emu(5181600), Emu(3900000))
    sh.name = '圖片位置'
    sh.fill.background()
    sh.line.color.rgb = RGBColor(0x99, 0x99, 0x99)
    sh.line.dash_style = MSO_LINE.DASH
    sh.line.width = Pt(2)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.text = '（在這裡放圖）\n' + hint
    for para in tf.paragraphs:
        for r in para.runs:
            r.font.size = Pt(20)
            r.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

def scene(slide, sents, vocab, hint=None, keep_prefix=''):
    """sents: [(zh_trad, es, en)]; vocab: [(zh_trad, es_short, es_full, en)]"""
    set_vocab(slide, [(v[0], v[1]) for v in vocab])
    lines = ['问：要怎么说：']
    lines += [f'- {z} {e}' for z, e, _ in sents]
    lines.append('')
    lines += [f'{v[0]}：{v[2]} / {v[3]}' for v in vocab]
    lines.append('')
    lines += [f'- {z} {n}' for z, _, n in sents]
    set_notes(slide, '\n'.join(lines))
    if hint:
        photo_placeholder(slide, hint)


# =============== fill-in pages ===============
def fill_page(slide, lines, es, en, vocab_notes):
    body = shape(slide, '文字方塊 2')._element.txBody
    bank_sh = shape(slide, '文字方塊 3')
    paras = body.findall(A + 'p')
    tpl = copy.deepcopy(paras[0])
    zh_r = copy.deepcopy(tpl.findall(A + 'r')[0])
    bl_r = copy.deepcopy(tpl.findall(A + 'r')[1])
    for para in paras:
        body.remove(para)
    words = []
    for ln in lines:
        words += re.findall(r'\[(.+?)\]', ln)
        simp = S(re.sub(r'\[(.+?)\]', '◯', ln))
        para = copy.deepcopy(tpl)
        for r in para.findall(A + 'r'):
            para.remove(r)
        end = para.find(A + 'endParaRPr')
        segs = re.split('(◯)', simp)
        for seg in segs:
            if seg == '':
                continue
            r = copy.deepcopy(bl_r if seg == '◯' else zh_r)
            r.find(A + 't').text = '_______' if seg == '◯' else seg
            if end is not None:
                end.addprevious(r)
            else:
                para.append(r)
        body.append(para)
    order = list(words)
    random.Random(11).shuffle(order)
    bank = bank_sh._element.txBody.findall(A + 'p')[0].findall(A + 'r')[0]
    bank.find(A + 't').text = '  '.join(S(w) for w in order)
    full = [re.sub(r'\[(.+?)\]', r'\1', ln) for ln in lines]
    vn = '\n'.join(vocab_notes)
    set_notes(slide, vn + '\n\n' + '\n'.join(full) + '\n\n' + '\n'.join(es) + '\n\n' + '\n'.join(en))



# ---------- sort pages ----------
def sort_page(slide, T, es, en, seed):
    n = len(T)
    boxes = sorted([s for s in slide.shapes if s.has_text_frame and re.match(r'^\d\.', s.text_frame.text or '')],
                   key=lambda s: s.top)
    assert len(boxes) == n, (len(boxes), n)
    idx = list(range(n))
    r = random.Random(seed)
    while True:
        r.shuffle(idx)
        if idx != sorted(idx):
            break
    shown = [T[i] for i in idx]
    for k, b in enumerate(boxes):
        b.text_frame.paragraphs[0].runs[0].text = f'{k+1}.{S(shown[k])}'
    order = [shown.index(t) + 1 for t in T]
    ans = '正確順序（填數字）：' + ' → '.join(map(str, order))
    set_notes(slide, ans + '\n\n' + '\n'.join(T) + '\n\n' + '\n'.join(es) + '\n\n' + '\n'.join(en))


lookup = {}
for i, (sents, vocab, hint) in enumerate(T.SCENES):
    scene(sl[1 + i], sents, [(v[0], v[1], v[1], v[2]) for v in vocab], hint=hint)
    for v in vocab:
        lookup[v[0]] = (v[1], v[2])
def vn(lines):
    words = []
    for ln in lines:
        words += re.findall(r'\[(.+?)\]', ln)
    out = []
    for w in words:
        if w in lookup:
            out.append(f'{w}：{lookup[w][0]} / {lookup[w][1]}')
        else:
            print('MISSING vocab for', w)
            out.append(w)
    # group 3 per line
    return ['  ///  '.join(out[i:i+3]) for i in range(0, len(out), 3)]
for j, (lines, es, en) in enumerate(T.FILL):
    fill_page(sl[7 + j], lines, es, en, vn(lines))
for j, (Tl, es, en, seed) in enumerate(T.SORT):
    sort_page(sl[9 + j], Tl, es, en, seed)
last = sl[13]
paras = shape(last, '文字方塊 2')._element.txBody.findall(A + 'p')
rv = T.REVIEW
newtxt = [(S(rv[0]), '?'), (S(rv[1]), '?'), (S(rv[2]), '?'), (S(rv[3]), None), (S(rv[4]), None), (S(rv[5]), None)]
for para, (a, b_) in zip(paras, newtxt):
    runs = para.findall(A + 'r')
    runs[0].find(A + 't').text = a
    if len(runs) > 1:
        if b_ is None:
            para.remove(runs[1])
        else:
            runs[1].find(A + 't').text = b_
set_notes(last, '角色扮演提示：' + ' / '.join(rv[3:6]) + '\n' + '\n'.join(rv[6]))
set_notes(sl[0], '接下來要學在' + T.OPEN + '可以用到的句子，有問題就要直接問，如果有你想到的其它的狀況也要問你，之後要靠這個生存')
for i, s in enumerate(sl):
    if i != 13 and 'show' in s._element.attrib:
        del s._element.attrib['show']
OUT = sys.argv[2]; os.makedirs(OUT, exist_ok=True)
p.save(os.path.join(OUT, T.FILE + '.pptx'))
print('saved', T.FILE)

# ---- lint ----
probs = []
if len(T.SCENES) != 6: probs.append('SCENES must be 6')
for i, (sents, vocab, hint) in enumerate(T.SCENES, 1):
    if not 3 <= len(sents) <= 4: probs.append(f'scene{i}: sentences {len(sents)} (need 3-4)')
    if not 5 <= len(vocab) <= 7: probs.append(f'scene{i}: vocab {len(vocab)} (need 5-7)')
    for v in vocab:
        if len(v[1]) > 22: probs.append(f'scene{i}: es gloss too long: {v[1]}')
if [len(x[0]) for x in T.SORT] != [6, 6, 8]: probs.append('SORT line counts must be 6,6,8')
for j, (lines, es, en) in enumerate(T.FILL, 1):
    if not 6 <= len(lines) <= 7 or len(es) != len(lines) or len(en) != len(lines): probs.append(f'fill{j}: line counts')
    if any(len(re.sub(r'[\[\]]', '', l)) > 25 for l in lines): probs.append(f'fill{j}: line >25 chars')
for j, (Tl, es, en, seed) in enumerate(T.SORT, 1):
    if len(es) != len(Tl) or len(en) != len(Tl): probs.append(f'sort{j}: line counts')
    if any(len(l) > 26 for l in Tl): probs.append(f'sort{j}: line >26 chars')
print('PROBLEMS:', 'none' if not probs else '\n  ' + '\n  '.join(probs))
